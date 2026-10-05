import os
import time
import json
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document


# 1. SETUP

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("api key not found")

client = Groq(api_key=my_api_key)
model = "openai/gpt-oss-120b"


# 2. JOB DESCRIPTION

job_description = f"""
Description
Do you want to solve real customer problems through innovative technology?
Do you enjoy working on scalable services in a collaborative team environment?
Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf
of our customers. Customer obsession is part of our company DNA.

Our Software Development Engineers (SDEs) use modern technology to solve
complex problems while seeing their work impact millions of customers.

Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary teams.
• Design and develop scalable solutions using cloud-native architectures and microservices.
• Participate in code reviews and technical documentation.
• Build and maintain resilient distributed systems.
• Leverage GenAI and AI-powered tools.
• Write clean, maintainable code.
• Work in an agile environment practicing CI/CD.
• Monitor, troubleshoot, and resolve production issues.

Basic Qualifications
- Experience with at least one general-purpose programming language such as
  Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structures, algorithms, and/or OOP
- Currently pursuing or obtaining a bachelor's degree in a related STEM field
- Must be 18 years of age or older

Preferred Qualifications
- Previous technical internship or demonstrated project experience
- AI tools for development productivity
- Cloud platforms, preferably AWS
- SQL and NoSQL databases
- Open-source contribution
- Version control systems
- Debugging and troubleshooting
- Ability to learn new technologies quickly
- Basic understanding of SDLC
- Strong problem-solving and analytical skills
- Good written and verbal communication skills
"""


# 3. PYDANTIC SCHEMA FOR JOB

class JobDescription(BaseModel):
    role: str
    req_skills: list[str]
    pref_skills: list[str]
    min_exp: float | None
    educ_req: list[str]
    responsibilities: list[str]

jobD_schema = JobDescription.model_json_schema()


# 4. CHAIN 1
# JOB DESCRIPTION → STRUCTURED REQUIREMENTS

system_prompt = f"""
You are an expert HR assistant.

Analyze the job description and extract structured information.

Return ONLY valid JSON matching this schema:

{jobD_schema}

IMPORTANT:
- Do NOT return the schema itself.
- Do NOT return fields like properties, title, or type.
- Fill the schema with actual information from the job description.
- If minimum experience is not mentioned, return null.
- If information for a list is missing, return an empty list.
- Do not invent information.
"""

user_prompt = f"""
Analyze the following job description:

{job_description}
"""


messages = [
    {
        "role": "system",
        "content": system_prompt
    },
    {
        "role": "user",
        "content": user_prompt
    }
]

response_format = {
    "type": "json_object"
}

response = client.chat.completions.create(
    model=model,
    messages=messages,
    response_format=response_format
)

raw_job = response.choices[0].message.content
job_data = json.loads(raw_job)
job = JobDescription(**job_data)

print("\n================ CHAIN 1 ================")
print("JOB REQUIREMENTS")
print(job.model_dump_json(indent=2))


# 5. PYDANTIC SCHEMAS FOR RESUME

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()


# 6. CHAIN 2
# JOB REQUIREMENTS + RESUME → STRUCTURED RESUME

def parse_resume(resume_text, job):

    system_prompt = f"""
You are an expert resume parser.

Extract structured information from the candidate resume.

The candidate is being evaluated for this job.

JOB REQUIREMENTS:

{job.model_dump_json(indent=2)}

Return ONLY valid JSON matching this schema:

{resume_schema}

Important rules:

1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
"""


    user_prompt = f"""
Parse the following resume:

{resume_text}
"""


    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]


    response_format = {
        "type": "json_object"
    }


    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )


    raw_output = response.choices[0].message.content

    data = json.loads(raw_output)

    resume = Resume(**data)

    return resume


# 7. CHAIN 3
# JOB + RESUME → MATCH SCORE

class MatchResult(BaseModel):
    score: float
    details: dict


def final_score(job, resume):

    match_schema = MatchResult.model_json_schema()

    prompt = f"""
You are an HR recruiter.

Compare the candidate's resume with the job requirements.

JOB REQUIREMENTS:

{job.model_dump_json(indent=2)}

CANDIDATE RESUME:

{resume.model_dump_json(indent=2)}

Return JSON matching this schema:

{match_schema}

Give me:

1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether experience requirement is met
5. Overall match percentage from 0 to 100
6. A short final verdict

Keep the response concise and easy to read.
"""


    message = {
        "role": "user",
        "content": prompt
    }


    messages = [message]


    response_format = {
        "type": "json_object"
    }


    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )


    data = json.loads(
        response.choices[0].message.content
    )

    return MatchResult(**data)


# 8. FUNCTIONS FOR READING RESUME

def read_pdf(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def read_docx(file_path):

    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"


    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_path):

    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


# 9. RUN THE PROMPT CHAIN

resume_folder = Path("resumes")
all_results = []


for file_path in resume_folder.iterdir():

    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue


    print("Processing:", file_path.name)

    # Read resume
    resume_text = read_resume(file_path)


    # CHAIN 2
    # Resume + Job Requirements

    parsed_resume = parse_resume(
        resume_text,
        job
    )

    print("\nCHAIN 2 - RESUME")
    print(
        parsed_resume.model_dump_json(
            indent=2
        )
    )


    time.sleep(5)
    # CHAIN 3
    # Job + Resume

    result = final_score(
        job,
        parsed_resume
    )


    print("\nCHAIN 3 - MATCH RESULT")
    print("Score:", result.score)
    print("Details:")
    print(result.details)


    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })


    time.sleep(5)


# 10. SORT RESULTS

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)


top_2 = all_results[:2]
worst_2 = all_results[-2:]

print("TOP 2 CANDIDATES")

for candidate in top_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])


print("LOWEST 2 CANDIDATES")
for candidate in worst_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])